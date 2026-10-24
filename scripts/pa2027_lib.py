# -*- coding: utf-8 -*-
"""Llibreria per generar els models del Programa d'actuacions 2027 a partir del Pla estratègic
i del Programa anual 2026 de cada espai. Ús: scripts/generar_model_2027.py"""
import re, warnings
import openpyxl
warnings.filterwarnings('ignore')

CODI_OO = re.compile(r'^\d{1,2}\.\d{1,2}\.\d{1,2}$')
EIXOS = [  # (codi 2 lletres, nom llista, OG inicial, OG final)
    ('NA', '1. Natura', 1, 3), ('CU', '2. Cultura', 4, 5), ('CO', '3. Coneixement', 6, 6),
    ('DE', '4. Desenvolupament local', 7, 9), ('US', '5. Ús públic i social', 10, 12),
    ('ED', '6. Educació, voluntariat i comunicació', 13, 14), ('GE', '7. Gestió i administració', 15, 16)]

def norm(v):
    return re.sub(r'\s+', ' ', str(v if v is not None else '')).strip()

def cell_text(v):
    """Text net que conserva els salts de línia com a espai."""
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return norm(v)

def codi_oo(v):
    """Normalitza un codi d'objectiu operatiu: '6.1.1' -> '06.1.1'; '0.2.2.2' no es toca."""
    s = norm(v)
    m = re.fullmatch(r'(\d{1,2})\.(\d{1,2})\.(\d{1,2})', s)
    return f'{int(m.group(1)):02d}.{m.group(2)}.{m.group(3)}' if m else s

def clau_oo(c):
    try: return tuple(int(x) for x in c.split('.'))
    except Exception: return (99, 99, 99)

def eix_de_oo(c):
    try: og = int(c.split('.')[0])
    except Exception: return None
    for cod, nom, a, b in EIXOS:
        if a <= og <= b: return cod, nom
    return None

def troba_capcalera(ws, maxfila=12):
    for r in range(1, maxfila + 1):
        vals = [norm(ws.cell(r, c).value).lower() for c in range(1, min(ws.max_column, 40) + 1)]
        if any(v.startswith('eixos') for v in vals) and any('operatius' in v or 'operatiu' in v for v in vals):
            return r
    return None

def llegeix_pla(path, fulla='PLA', cap=None, mapeig=None):
    """Retorna una llista d'accions del Pla (una fila per acció; un objectiu sense acció dóna una fila amb accio='')."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[fulla]
    hr = cap or troba_capcalera(ws)
    H = {c: norm(ws.cell(hr, c).value).lower() for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value is not None}
    def col(*pref, excl=()):
        for c, h in H.items():
            if any(h.startswith(p) for p in pref) and not any(e in h for e in excl): return c
        return None
    M = dict(og=3, oe=5, op=6, optext=7, comu=col('comú'), oblig=col('obligatori'),
             concret=col("objectiu operatiu concret", "redactat de l'op", "objectius operatius pnzvg", "objectiu operatiu alta", "objectius operatius alta"),
             incorp=col('objectiu amb mesures', 'objectiu aplicable'), prior=col('prioritat'),
             accio=col("propostes genèriques", "propostes d'accions", "proposta concreta", "proposta d'accions", 'accions', 'accions (afegit', "proposta d'accions pla"),
             ind=col('indicadors', 'indicador'), mn=col('mínima', 'mínim', 'valor mínim'), mx=col('màxim', 'òptim', 'valor objectiu'),
             area=col('àrea responsable'), pers=col('persona responsable'))
    if mapeig: M.update(mapeig)
    out = []; cur_og = cur_oe = ''; cur = None
    for r in range(hr + 1, ws.max_row + 1):
        g = lambda k: cell_text(ws.cell(r, M[k]).value) if M.get(k) else ''
        if g('og'): cur_og = g('og')
        if g('oe'): cur_oe = g('oe')
        op = norm(ws.cell(r, M['op']).value)
        if CODI_OO.match(op):
            cur = dict(oo=codi_oo(op), optext=g('optext'), concret=g('concret'), comu=g('comu'), oblig=g('oblig'),
                       incorp=g('incorp'), og=cur_og, oe=cur_oe)
            out.append(dict(cur, accio='', ind='', mn='', mx='', prior=g('prior'), area='', pers='', _buit=True))
        if cur is None: continue
        accio = g('accio')
        if accio:
            fila = dict(cur, accio=accio, ind=g('ind'), mn=cell_text(ws.cell(r, M['mn']).value) if M.get('mn') else '',
                        mx=cell_text(ws.cell(r, M['mx']).value) if M.get('mx') else '', prior=g('prior'),
                        area=g('area'), pers=g('pers'), _buit=False)
            if out and out[-1]['oo'] == cur['oo'] and out[-1]['_buit'] and out[-1] is not fila:
                out.pop()  # la fila de l'objectiu sense acció queda substituïda per la primera acció
            out.append(fila)
    out.sort(key=lambda x: clau_oo(x['oo']))  # estable: manté l'ordre intern de cada objectiu
    return out

def llegeix_pa2026(path, fulla='PROGRAMA ANUAL', cap=3):
    wb = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path)
    ws = wb[fulla]
    H = {c: norm(ws.cell(cap, c).value).lower() for c in range(1, ws.max_column + 1) if ws.cell(cap, c).value is not None}
    def col(*pref, excl=()):
        for c, h in H.items():
            if any(h.startswith(p) for p in pref) and not any(e in h for e in excl): return c
    C = dict(codi=1, rae=col('codi rae'), accio=col('actuació'), oo=col("codi de l'objectiu", 'codi oo'),
             rel=col("codis d'altres"), eix=col('eix'), desp=col("tipus d'actuació segons model", 'contracte menor'),
             part=col('partida'), mapt=col("tipus d'actuació segons map"), imp=col('import'), grau=col('grau'),
             obs=col('observacions'), gest=col('gestor'), rec=col('actuació recurrent'))
    files = []
    for r in range(cap + 1, ws.max_row + 1):
        if not ws.cell(r, 1).value: continue
        g = lambda k: ws.cell(r, C[k]).value if C.get(k) else None
        imp = g('imp'); f = wf[fulla].cell(r, C['imp']).value if C.get('imp') else None
        obs = cell_text(g('obs'))
        if isinstance(f, str) and f.startswith('='):
            nota = f"Import 2026 calculat amb la fórmula {f}"
            obs = f'{obs}. {nota}' if obs else nota
        rel = [codi_oo(x) for x in re.split(r'[;,]\s*', norm(g('rel'))) if x.strip()]
        codi = norm(ws.cell(r, 1).value); codi = re.sub(r'^25_', '26_', codi)
        rae = g('rae'); rae = int(rae) if isinstance(rae, float) and rae.is_integer() else rae
        eix = norm(g('eix')).replace('1.Natura', '1. Natura').replace('Gestió i Administració', 'Gestió i administració')
        files.append(dict(codi=codi, rae=rae if rae is not None else '', accio=cell_text(g('accio')), oo=codi_oo(g('oo')),
                          rel=', '.join(rel), eix=eix, desp=cell_text(g('desp')), part=cell_text(g('part')),
                          mapt=cell_text(g('mapt')), imp=imp if imp is not None else 0, grau=g('grau'),
                          obs=obs, gest=cell_text(g('gest')), rec=cell_text(g('rec'))))
    return files


# ════════════════════════════════════════════════════════════════════
#  Construcció del model Excel 2027
# ════════════════════════════════════════════════════════════════════
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL

VERSIO_MODEL = '1.0'
FILES_PA = 700            # files preparades per al programa 2027
PRIMERA = 6               # primera fila de dades
TEAL = '088074'
F_CAP = PatternFill('solid', fgColor=TEAL); F_AUTO = PatternFill('solid', fgColor='EDEDED')
F_NOTA = PatternFill('solid', fgColor='FFF8E1'); F_SEC = PatternFill('solid', fgColor='E0F2F1')
BLANC = Font(bold=True, color='FFFFFF'); NEGRETA = Font(bold=True)
FINA = Side(style='thin', color='BDBDBD'); BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
WRAP = Alignment(wrap_text=True, vertical='top')

def num_o_text(v):
    s = norm(v)
    if s == '': return None
    try:
        f = float(s.replace(',', '.'))
        return int(f) if f.is_integer() else f
    except ValueError:
        return s

def llistes_base(path_base):
    wb = openpyxl.load_workbook(path_base, data_only=True); ws = wb['Llistes']
    col = lambda c, a, b: [norm(ws.cell(r, c).value) for r in range(a, b + 1) if norm(ws.cell(r, c).value)]
    despesa = col(2, 2, 15)
    if 'Mitjans propis' not in despesa:
        i = despesa.index('Altres') if 'Altres' in despesa else len(despesa)
        despesa.insert(i, 'Mitjans propis')
    mapt = [m.replace('conservación i restauración', 'conservació i restauració') for m in col(3, 2, 19)]
    partides = col(7, 3, 60)
    return despesa, mapt, partides

def gestors_parc(path_pa2026, fragment):
    wb = openpyxl.load_workbook(path_pa2026, data_only=True)
    if 'Gestors' not in wb.sheetnames: return []
    ws = wb['Gestors']; frag = fragment.lower()
    for c in range(1, ws.max_column + 1):
        if frag in norm(ws.cell(1, c).value).lower():
            return sorted({norm(ws.cell(r, c).value) for r in range(2, ws.max_row + 1) if norm(ws.cell(r, c).value)})
    return []

def capcalera(ws, fila, textos, amples=None, auto=()):
    for i, t in enumerate(textos, 1):
        c = ws.cell(fila, i, t); c.fill = F_CAP; c.font = BLANC; c.border = BORDA
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        if amples: ws.column_dimensions[CL(i)].width = amples[i - 1]
    ws.row_dimensions[fila].height = 48

def construeix_model(cfg, sortida, path_base):
    """cfg: codi, nom, pla (path), fulla_pla, pa2026 (path), gestors_fragment, cap_pla (opc.), mapeig_pla (opc.)"""
    pla = llegeix_pla(cfg['pla'], cfg.get('fulla_pla', 'PLA'), cfg.get('cap_pla'), cfg.get('mapeig_pla'))
    pa26 = llegeix_pa2026(cfg['pa2026'], cfg.get('fulla_pa', 'PROGRAMA ANUAL'), cfg.get('cap_pa', 3)) if cfg.get('pa2026') else []
    despesa, mapt, partides = llistes_base(path_base)
    gestors = gestors_parc(cfg['pa2026'], cfg['gestors_fragment']) if cfg.get('pa2026') else []
    oos = []; vist = set()
    for x in pla:
        if x['oo'] not in vist:
            vist.add(x['oo']); oos.append(x)
    wb = Workbook()

    # ── Instruccions ────────────────────────────────────────────────
    wi = wb.active; wi.title = 'Instruccions'
    wi.column_dimensions['A'].width = 34; wi.column_dimensions['B'].width = 110
    wi['A1'] = f"Programa d'actuacions 2027 — {cfg['nom']}"; wi['A1'].font = Font(bold=True, size=16, color=TEAL)
    wi['A2'] = 'Versió del model:'; wi['B2'] = VERSIO_MODEL
    wi['A3'] = 'Codi de la unitat al Motor:'; wi['B3'] = cfg['codi']
    wi['A4'] = 'Pla estratègic de referència:'; wi['B4'] = cfg.get('pla_nom', '')
    for r in (2, 3, 4): wi.cell(r, 1).font = NEGRETA
    pas = [
        ('COM OMPLIR-LO', ''),
        ('1. Mireu el Pla', "La pestanya PLA ESTRATÈGIC porta el pla aprovat del vostre espai (objectius, accions, indicadors i fites). No s'ha de modificar. La pestanya PA 2026 (consulta) porta les actuacions del 2026, per si en voleu repetir o adaptar alguna."),
        ('2. Ompliu PROGRAMA ANUAL 2027', "Una actuació per fila. Les columnes blanques són per omplir i les grises es calculen soles. Gairebé tot són desplegables: no cal escriure codis a mà."),
        ('3. Mireu la columna «Revisió»', "Cada fila mostra ✔ si està completa o diu què hi falta (import, partida RAE, gestor, codi RAE repetit…). Abans d'enviar el fitxer, no hauria de quedar cap avís."),
        ('4. Envieu el fitxer', "El fitxer s'ha de poder convertir al JSON del Motor. No canvieu els noms de les pestanyes ni les columnes."),
        ('', ''),
        ('COLUMNES DE PROGRAMA ANUAL 2027', ''),
        ('Codi (automàtic)', "Es genera sol: any 27, eix i número correlatiu (27_NA_01). Depèn de l'ordre de les files; no les reordeneu ni n'esborreu cap un cop enviat el fitxer. Per anul·lar una actuació, buideu-ne les cel·les."),
        ('Codi RAE', "Només si l'actuació té codi RAE. Ha de ser únic."),
        ('Actuació', "Què farà el parc el 2027. Ha de derivar d'una acció del Pla estratègic."),
        ('Objectiu operatiu principal', "Desplegable amb els objectius del vostre Pla. El nom de l'objectiu i l'eix s'omplen sols."),
        ('Altres objectius relacionats', "Fins a tres, amb desplegable. Opcional."),
        ('Tipus de despesa', "Desplegable. «Mitjans propis» és per a les actuacions sense contractació externa."),
        ('Partida RAE', "Només per a «Contracte menor a càrrec del RAE». Obligatòria en aquest cas."),
        ('Tipus d'+"actuació segons MAP", "Desplegable amb les tipologies del Marc d'Acció Prioritària. Si cap no hi encaixa, «Altres»."),
        ('Import (€)', "Import previst. Si no té cost directe, poseu 0. No s'admeten valors negatius."),
        ('Acció i indicador del Pla', "Desplegable amb les accions del Pla per a l'objectiu triat. La fita mínima i la màxima s'omplen soles."),
        ('Grau de compliment', "Percentatge (0–100%) de la fita màxima que aconseguirà l'actuació. Si és qualitatiu, expliqueu-ho a Observacions."),
        ('Gestor responsable', "Desplegable amb l'equip del parc. Si hi falta algú, es pot escriure el nom."),
        ('Actuació recurrent', "Sí si es repeteix cada any."),
        ('Naturalesa del contracte', "Opcional: Obra, Servei o Subministrament. Ajuda el Motor a proposar el procediment."),
        ('', ''),
        ('QUÈ HA CANVIAT RESPECTE AL MODEL ANTERIOR', ''),
        ('', "Es porta el Pla estratègic de l'espai amb indicadors i fites; l'objectiu i l'eix es trien i es calculen sols; el codi és automàtic; hi ha validacions i una columna de revisió; s'hi han afegit el gestor, l'actuació recurrent i el tipus de despesa «Mitjans propis»."),
    ]
    r = 6
    for a, b in pas:
        wi.cell(r, 1, a).font = NEGRETA; wi.cell(r, 2, b).alignment = WRAP; wi.cell(r, 1).alignment = WRAP
        if a.isupper() and not b:
            wi.cell(r, 1).fill = F_SEC; wi.cell(r, 2).fill = F_SEC
        r += 1

    # ── PLA ESTRATÈGIC (fulla de referència) ────────────────────────
    wp = wb.create_sheet('PLA ESTRATÈGIC')
    wp['A1'] = f"PLA ESTRATÈGIC — {cfg['nom']} (només consulta: no modificar)"; wp['A1'].font = Font(bold=True, size=13, color=TEAL)
    cap = ['Eix', 'Codi objectiu operatiu', 'Objectiu operatiu (estàndard)', "Objectiu concret a l'espai", 'Acció del Pla', 'Indicador', 'Fita mínima',
           'Fita màxima', 'Prioritat', 'Àrea responsable', 'Persona responsable', 'Objectiu general', 'Objectiu estratègic', 'Comú / Específic',
           'Obligatori / Optatiu', 'Incorporable a la planificació anual', 'Clau (per als desplegables)']
    amp = [22, 12, 45, 45, 60, 40, 10, 10, 11, 18, 18, 40, 40, 12, 12, 14, 70]
    capcalera(wp, 3, cap, amp)
    claus = set(); fila = 4
    for x in pla:
        e = eix_de_oo(x['oo']); eix = e[1] if e else ''
        base = f"{x['oo']} · {(x['accio'] or '(cap acció definida al Pla)')[:150]}"
        if x['ind']: base += f" — {x['ind'][:100]}"
        clau = base; n = 2
        while clau in claus: clau = f'{base} #{n}'; n += 1
        claus.add(clau)
        vals = [eix, x['oo'], x['optext'], x['concret'] or x['optext'], x['accio'], x['ind'], num_o_text(x['mn']), num_o_text(x['mx']),
                x['prior'], x['area'], x['pers'], x['og'], x['oe'], x['comu'], x['oblig'], x['incorp'], clau]
        for i, v in enumerate(vals, 1):
            c = wp.cell(fila, i, v); c.alignment = WRAP; c.border = BORDA
        fila += 1
    ULT_PLA = fila - 1
    wp.freeze_panes = 'C4'; wp.auto_filter.ref = f'A3:{CL(len(cap))}{ULT_PLA}'

    # ── Objectius (llista única per als desplegables) ───────────────
    wo = wb.create_sheet('Objectius')
    capcalera(wo, 1, ['Codi', 'Objectiu operatiu (estàndard)', "Objectiu concret a l'espai", 'Eix', 'Codi eix', 'Objectiu general', 'Objectiu estratègic'], [10, 50, 50, 28, 9, 40, 40])
    for i, x in enumerate(oos, 2):
        e = eix_de_oo(x['oo']) or ('', '')
        for j, v in enumerate([x['oo'], x['optext'], x['concret'] or x['optext'], e[1], e[0], x['og'], x['oe']], 1):
            wo.cell(i, j, v)
    N_OO = len(oos) + 1

    # ── Llistes ─────────────────────────────────────────────────────
    wl = wb.create_sheet('Llistes')
    capcalera(wl, 1, ['Tipus de despesa', 'Tipologia MAP', 'Eix', 'Codi eix', 'Partides RAE', 'Naturalesa', 'Sí/No'], [38, 70, 36, 9, 60, 18, 8])
    for i, v in enumerate(despesa, 2): wl.cell(i, 1, v)
    for i, v in enumerate(mapt, 2): wl.cell(i, 2, v)
    for i, (cod, nom, a, b) in enumerate(EIXOS, 2): wl.cell(i, 3, nom); wl.cell(i, 4, cod)
    for i, v in enumerate(partides, 2): wl.cell(i, 5, v)
    for i, v in enumerate(['Obra', 'Servei', 'Subministrament'], 2): wl.cell(i, 6, v)
    for i, v in enumerate(['Sí', 'No'], 2): wl.cell(i, 7, v)

    # ── Gestors ─────────────────────────────────────────────────────
    wg = wb.create_sheet('Gestors')
    capcalera(wg, 1, ["Gestors de l'espai"], [40])
    for i, v in enumerate(gestors or [''], 2): wg.cell(i, 1, v)
    N_GEST = max(len(gestors), 1) + 1

    # ── PROGRAMA ANUAL 2027 ─────────────────────────────────────────
    ws = wb.create_sheet('PROGRAMA ANUAL 2027', 1)
    ws['A1'] = f"PROGRAMA D'ACTUACIONS 2027 — {cfg['nom']}"; ws['A1'].font = Font(bold=True, size=14, color=TEAL)
    ws['A2'] = 'Unitat:'; ws['B2'] = cfg['codi']; ws['A3'] = 'Model:'; ws['B3'] = f'v{VERSIO_MODEL}'
    ws['A2'].font = ws['A3'].font = NEGRETA
    ws['D2'] = 'Blanc = per omplir · Gris = automàtic · Reviseu la columna «Revisió» abans d\'enviar el fitxer.'; ws['D2'].fill = F_NOTA
    caps = ['Codi (automàtic)', 'Codi RAE', 'Actuació', 'Objectiu operatiu principal', "Objectiu concret a l'espai (automàtic)", 'Altre objectiu relacionat 1',
            'Altre objectiu relacionat 2', 'Altre objectiu relacionat 3', 'Eix (automàtic)', 'Tipus de despesa', 'Partida RAE', "Tipus d'actuació segons MAP",
            'Import (€)', 'Acció i indicador del Pla', 'Fita mínima (automàtic)', 'Fita màxima (automàtic)', 'Grau de compliment (% de la fita)',
            'Gestor responsable', 'Actuació recurrent', 'Naturalesa del contracte', 'Observacions', 'Revisió']
    ampl = [13, 10, 55, 13, 40, 11, 11, 11, 24, 30, 36, 40, 13, 60, 11, 11, 13, 24, 11, 16, 36, 40]
    capcalera(ws, 5, caps, ampl)
    ult = PRIMERA + FILES_PA - 1
    O = f'Objectius!$A$2:$A${N_OO}'
    for r in range(PRIMERA, ult + 1):
        ws[f'A{r}'] = f'=IF($D{r}="","","27_"&IFERROR(INDEX(Objectius!$E$2:$E${N_OO},MATCH($D{r},{O},0)),"??")&"_"&TEXT(COUNTIF($I${PRIMERA}:$I{r},$I{r}),"00"))'
        ws[f'E{r}'] = f'=IF($D{r}="","",IFERROR(INDEX(Objectius!$C$2:$C${N_OO},MATCH($D{r},{O},0)),"Objectiu no trobat al Pla"))'
        ws[f'I{r}'] = f'=IF($D{r}="","",IFERROR(INDEX(Objectius!$D$2:$D${N_OO},MATCH($D{r},{O},0)),""))'
        ws[f'O{r}'] = f"=IF($N{r}=\"\",\"\",IFERROR(INDEX('PLA ESTRATÈGIC'!$G$4:$G${ULT_PLA},MATCH($N{r},'PLA ESTRATÈGIC'!$Q$4:$Q${ULT_PLA},0)),\"\"))"
        ws[f'P{r}'] = f"=IF($N{r}=\"\",\"\",IFERROR(INDEX('PLA ESTRATÈGIC'!$H$4:$H${ULT_PLA},MATCH($N{r},'PLA ESTRATÈGIC'!$Q$4:$Q${ULT_PLA},0)),\"\"))"
        ws[f'V{r}'] = (f'=IF(COUNTA($C{r},$D{r},$J{r},$M{r})=0,"",IF(TRIM('
            f'IF($C{r}="","Falta l\'actuació. ","")&IF($D{r}="","Falta l\'objectiu. ","")&IF($J{r}="","Falta el tipus de despesa. ","")'
            f'&IF(AND($J{r}<>"",$J{r}<>"Mitjans propis",$J{r}<>"Altres",$M{r}=""),"Falta l\'import. ","")'
            f'&IF(AND($J{r}="Contracte menor a càrrec del RAE",$K{r}=""),"Falta la partida RAE. ","")'
            f'&IF(AND($B{r}<>"",COUNTIF($B${PRIMERA}:$B${ult},$B{r})>1),"Codi RAE repetit. ","")'
            f'&IF(AND($B{r}<>"",$J{r}<>"",$J{r}<>"Contracte menor a càrrec del RAE"),"Té codi RAE però no és de tipus RAE. ","")'
            f'&IF($N{r}="","Falta l\'indicador del Pla. ","")&IF($R{r}="","Falta el gestor. ","")&IF($S{r}="","Falta indicar si és recurrent. ",""))="","✔",'
            f'TRIM(IF($C{r}="","Falta l\'actuació. ","")&IF($D{r}="","Falta l\'objectiu. ","")&IF($J{r}="","Falta el tipus de despesa. ","")'
            f'&IF(AND($J{r}<>"",$J{r}<>"Mitjans propis",$J{r}<>"Altres",$M{r}=""),"Falta l\'import. ","")'
            f'&IF(AND($J{r}="Contracte menor a càrrec del RAE",$K{r}=""),"Falta la partida RAE. ","")'
            f'&IF(AND($B{r}<>"",COUNTIF($B${PRIMERA}:$B${ult},$B{r})>1),"Codi RAE repetit. ","")'
            f'&IF(AND($B{r}<>"",$J{r}<>"",$J{r}<>"Contracte menor a càrrec del RAE"),"Té codi RAE però no és de tipus RAE. ","")'
            f'&IF($N{r}="","Falta l\'indicador del Pla. ","")&IF($R{r}="","Falta el gestor. ","")&IF($S{r}="","Falta indicar si és recurrent. ",""))))')
        for c in range(1, 23):
            cel = ws.cell(r, c); cel.border = BORDA; cel.alignment = WRAP
            if c in (1, 5, 9, 15, 16, 22): cel.fill = F_AUTO
        ws[f'M{r}'].number_format = '#,##0.00 "€"'; ws[f'Q{r}'].number_format = '0%'
    # validacions
    def dv(rang, formula, tipus='list', **kw):
        d = DataValidation(type=tipus, formula1=formula, allow_blank=True, **kw)
        d.error = kw.get('error', 'Valor no admès. Trieu un valor de la llista.'); d.errorTitle = 'Valor no admès'
        d.showErrorMessage = True; ws.add_data_validation(d); d.add(rang); return d
    dv(f'D{PRIMERA}:D{ult}', f'={O}')
    for c in 'FGH': dv(f'{c}{PRIMERA}:{c}{ult}', f'={O}')
    dv(f'J{PRIMERA}:J{ult}', f'=Llistes!$A$2:$A${len(despesa)+1}')
    dv(f'K{PRIMERA}:K{ult}', f'=Llistes!$E$2:$E${len(partides)+1}')
    dv(f'L{PRIMERA}:L{ult}', f'=Llistes!$B$2:$B${len(mapt)+1}')
    d = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    d.error = "L'import ha de ser un número igual o superior a 0."; d.errorTitle = 'Import no vàlid'; d.showErrorMessage = True
    ws.add_data_validation(d); d.add(f'M{PRIMERA}:M{ult}')
    dn = DataValidation(type='list', formula1=f"=OFFSET('PLA ESTRATÈGIC'!$Q$1,MATCH($D{PRIMERA},'PLA ESTRATÈGIC'!$B:$B,0)-1,0,COUNTIF('PLA ESTRATÈGIC'!$B:$B,$D{PRIMERA}),1)", allow_blank=True)
    dn.error = "Trieu una acció del Pla estratègic per a l'objectiu seleccionat."; dn.errorTitle = 'Acció no vàlida'; dn.showErrorMessage = True
    ws.add_data_validation(dn); dn.add(f'N{PRIMERA}:N{ult}')
    dg = DataValidation(type='decimal', operator='between', formula1='0', formula2='1', allow_blank=True)
    dg.error = 'Poseu un percentatge entre 0% i 100%.'; dg.errorTitle = 'Percentatge no vàlid'; dg.showErrorMessage = True
    ws.add_data_validation(dg); dg.add(f'Q{PRIMERA}:Q{ult}')
    dge = DataValidation(type='list', formula1=f'=Gestors!$A$2:$A${N_GEST}', allow_blank=True, errorStyle='warning')
    dge.error = "Aquest nom no és a la llista de gestors. Voleu continuar igualment?"; dge.errorTitle = 'Gestor nou'; dge.showErrorMessage = True
    ws.add_data_validation(dge); dge.add(f'R{PRIMERA}:R{ult}')
    dv(f'S{PRIMERA}:S{ult}', '=Llistes!$G$2:$G$3')
    dv(f'T{PRIMERA}:T{ult}', '=Llistes!$F$2:$F$4')
    ws.conditional_formatting.add(f'V{PRIMERA}:V{ult}', FormulaRule(formula=[f'AND($V{PRIMERA}<>"",$V{PRIMERA}<>"✔")'], fill=PatternFill('solid', bgColor='FFCDD2')))
    ws.conditional_formatting.add(f'V{PRIMERA}:V{ult}', FormulaRule(formula=[f'$V{PRIMERA}="✔"'], fill=PatternFill('solid', bgColor='C8E6C9')))
    ws.freeze_panes = 'D6'
    ws.sheet_properties.tabColor = TEAL

    # ── PA 2026 (consulta) ──────────────────────────────────────────
    w6 = wb.create_sheet('PA 2026 (consulta)')
    w6['A1'] = f"PROGRAMA D'ACTUACIONS 2026 — {cfg['nom']} (només consulta)"; w6['A1'].font = Font(bold=True, size=13, color=TEAL)
    w6['A2'] = "Codis amb prefix 26_, objectius i eixos normalitzats. Si no hi ha actuacions, el parc no tenia programa 2026 al Motor."
    cap6 = ['Codi 2026', 'Codi RAE', 'Actuació', 'Objectiu operatiu', 'Altres objectius', 'Eix', 'Tipus de despesa', 'Partida RAE', 'Tipus MAP', 'Import (€)',
            'Grau de compliment', 'Observacions', 'Gestor', 'Recurrent', "L'objectiu és al Pla 2027?"]
    capcalera(w6, 4, cap6, [12, 10, 60, 11, 18, 22, 30, 34, 38, 12, 11, 36, 22, 10, 14])
    for i, p in enumerate(pa26, 5):
        v = [p['codi'], p['rae'], p['accio'], p['oo'], p['rel'], p['eix'], p['desp'], p['part'], p['mapt'], p['imp'], p['grau'], p['obs'], p['gest'], p['rec']]
        for j, x in enumerate(v, 1):
            c = w6.cell(i, j, x); c.alignment = WRAP; c.border = BORDA
        w6.cell(i, 15, f'=IF(COUNTIF(Objectius!$A$2:$A${N_OO},D{i})>0,"Sí","No")').border = BORDA
        w6.cell(i, 10).number_format = '#,##0.00 "€"'; w6.cell(i, 11).number_format = '0%'
    w6.freeze_panes = 'D5'
    if pa26: w6.auto_filter.ref = f'A4:O{4+len(pa26)}'

    # ordre de pestanyes i colors
    ordre = ['Instruccions', 'PROGRAMA ANUAL 2027', 'PLA ESTRATÈGIC', 'PA 2026 (consulta)', 'Llistes', 'Objectius', 'Gestors']
    wb._sheets = [wb[n] for n in ordre]
    for n in ('Llistes', 'Objectius', 'Gestors'): wb[n].sheet_properties.tabColor = '9E9E9E'
    wb['PLA ESTRATÈGIC'].sheet_properties.tabColor = '5C6BC0'; wb['PA 2026 (consulta)'].sheet_properties.tabColor = '5C6BC0'
    wb.active = 1
    wb.save(sortida)
    return dict(oo=len(oos), accions=sum(1 for x in pla if x['accio']), amb_ind=sum(1 for x in pla if x['ind']), pa26=len(pa26), gestors=len(gestors))
